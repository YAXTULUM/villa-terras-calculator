"""Five-tier LP/GP real estate distribution waterfall.

Run: python waterfall.py --input example_deal.json --output-dir output
"""
from __future__ import annotations

import argparse
import json
import math
from dataclasses import dataclass
from datetime import date
from pathlib import Path


@dataclass(frozen=True)
class Terms:
    preferred_rate: float = 0.08
    catchup_enabled: bool = True
    catchup_share: float = 0.20
    tier3_irr: float = 0.12
    tier3_multiple: float = 1.5
    tier3_lp_share: float = 0.65
    tier4_irr: float = 0.15
    tier4_multiple: float = 2.0
    tier4_lp_share: float = 0.55
    residual_lp_share: float = 0.50

    def validate(self):
        rates = ("preferred_rate", "tier3_irr", "tier4_irr")
        shares = ("tier3_lp_share", "tier4_lp_share", "residual_lp_share")
        for name in rates:
            value = getattr(self, name)
            if not math.isfinite(value) or value < 0:
                raise ValueError(f"Invalid {name}: {value}")
        for name in shares:
            value = getattr(self, name)
            if not math.isfinite(value) or not 0 <= value <= 1:
                raise ValueError(f"Invalid {name}: {value}")
        if not math.isfinite(self.catchup_share) or not 0 <= self.catchup_share < 1:
            raise ValueError(f"Invalid catchup_share: {self.catchup_share}")
        if not self.tier3_lp_share or not self.tier4_lp_share:
            raise ValueError("Hurdle tier LP shares must be positive")
        if self.tier3_irr > self.tier4_irr or self.tier3_multiple > self.tier4_multiple:
            raise ValueError("Hurdles must increase across tiers")
        if (not math.isfinite(self.tier3_multiple) or not math.isfinite(self.tier4_multiple)
                or self.tier3_multiple < 1 or self.tier4_multiple < 1):
            raise ValueError("Equity multiples must be at least 1.0")


def xirr(flows: list[tuple[date, float]]) -> float | None:
    if not flows or not any(v < 0 for _, v in flows) or not any(v > 0 for _, v in flows):
        return None
    start = flows[0][0]

    def npv(rate):
        return sum(amount / (1 + rate) ** ((day - start).days / 365)
                   for day, amount in flows)

    low, high = -0.9999, 1.0
    while npv(low) * npv(high) > 0 and high < 1e8:
        high = 2 * high + 1
    if npv(low) * npv(high) > 0:
        return None
    for _ in range(150):
        mid = (low + high) / 2
        if npv(mid) > 0:
            low = mid
        else:
            high = mid
    return (low + high) / 2


def calculate(periods: list[dict], terms: Terms = Terms()) -> dict:
    terms.validate()
    if not periods:
        raise ValueError("At least one dated period is required")
    rows = sorted(periods, key=lambda p: p["date"])
    if len({p["date"] for p in rows}) != len(rows):
        raise ValueError("Period dates must be unique")
    capital_lp = capital_gp = preferred_unpaid = 0.0
    lp_contributed = lp_paid_total = preferred_paid_total = catchup_paid_total = 0.0
    hurdle12 = hurdle15 = 0.0
    result = []
    lp_flows, gp_flows = [], []
    last_date = None
    for p in rows:
        day = date.fromisoformat(p["date"])
        lp_in, gp_in, cash = (float(p.get(k, 0)) for k in ("lp_contribution", "gp_contribution", "cash"))
        if any(not math.isfinite(value) or value < 0 for value in (lp_in, gp_in, cash)):
            raise ValueError("Contributions and distributable cash must be finite and nonnegative")
        elapsed = (day - last_date).days if last_date else 0
        if last_date:
            preferred_unpaid += capital_lp * terms.preferred_rate * elapsed / 365
            hurdle12 *= (1 + terms.tier3_irr) ** (elapsed / 365)
            hurdle15 *= (1 + terms.tier4_irr) ** (elapsed / 365)
        capital_lp += lp_in
        capital_gp += gp_in
        lp_contributed += lp_in
        hurdle12 += lp_in
        hurdle15 += lp_in

        # Capital repayment is proportionate to each party's unreturned balance.
        capital_cash = min(cash, capital_lp + capital_gp)
        lp_cap = capital_cash * capital_lp / (capital_lp + capital_gp) if capital_cash else 0.0
        gp_cap = capital_cash - lp_cap
        capital_lp -= lp_cap
        capital_gp -= gp_cap
        remaining = cash - capital_cash

        lp_pref = min(remaining, preferred_unpaid)
        preferred_unpaid -= lp_pref
        preferred_paid_total += lp_pref
        remaining -= lp_pref
        catchup_target = (terms.catchup_share / (1 - terms.catchup_share) * preferred_paid_total
                          if terms.catchup_enabled else 0.0)
        gp_catchup = min(remaining, max(0.0, catchup_target - catchup_paid_total))
        catchup_paid_total += gp_catchup
        remaining -= gp_catchup

        lp_before = lp_cap + lp_pref

        def hurdle_allocation(available, balance, multiple, share, paid_before):
            needed = max(0.0, balance - paid_before,
                         multiple * lp_contributed - lp_paid_total - paid_before)
            return min(available, needed / share)

        tier3_cash = hurdle_allocation(remaining, hurdle12, terms.tier3_multiple,
                                        terms.tier3_lp_share, lp_before)
        lp_tier3 = tier3_cash * terms.tier3_lp_share
        remaining -= tier3_cash
        tier4_cash = hurdle_allocation(remaining, hurdle15, terms.tier4_multiple,
                                        terms.tier4_lp_share, lp_before + lp_tier3)
        lp_tier4 = tier4_cash * terms.tier4_lp_share
        remaining -= tier4_cash
        residual_cash = max(0.0, remaining)
        lp_residual = residual_cash * terms.residual_lp_share
        lp_paid = lp_before + lp_tier3 + lp_tier4 + lp_residual
        gp_paid = cash - lp_paid
        hurdle12 -= lp_paid
        hurdle15 -= lp_paid
        lp_paid_total += lp_paid
        lp_flows.append((day, lp_paid - lp_in))
        gp_flows.append((day, gp_paid - gp_in))
        record = {"date": p["date"], "lp_contribution": lp_in, "gp_contribution": gp_in,
                  "cash": cash, "lp_capital": lp_cap, "gp_capital": gp_cap,
                  "lp_preferred": lp_pref, "gp_catchup": gp_catchup,
                  "lp_tier3": lp_tier3, "gp_tier3": tier3_cash - lp_tier3,
                  "lp_tier4": lp_tier4, "gp_tier4": tier4_cash - lp_tier4,
                  "lp_residual": lp_residual, "gp_residual": residual_cash - lp_residual,
                  "lp_paid": lp_paid, "gp_paid": gp_paid,
                  "lp_capital_outstanding": capital_lp, "gp_capital_outstanding": capital_gp,
                  "lp_preferred_outstanding": preferred_unpaid,
                  "cash_difference": cash - lp_paid - gp_paid}
        result.append(record)
        last_date = day
    total_lp_in = sum(r["lp_contribution"] for r in result)
    total_gp_in = sum(r["gp_contribution"] for r in result)
    total_lp = sum(r["lp_paid"] for r in result)
    total_gp = sum(r["gp_paid"] for r in result)
    return {"terms": vars(terms), "periods": result, "summary": {
        "lp_contributed": total_lp_in, "gp_contributed": total_gp_in,
        "lp_distributions": total_lp, "gp_distributions": total_gp,
        "lp_equity_multiple": total_lp / total_lp_in if total_lp_in else None,
        "gp_equity_multiple": total_gp / total_gp_in if total_gp_in else None,
        "lp_irr": xirr(lp_flows), "gp_irr": xirr(gp_flows),
        "cash_difference": sum(r["cash_difference"] for r in result),
        "lp_capital_outstanding": capital_lp, "gp_capital_outstanding": capital_gp,
        "lp_preferred_outstanding": preferred_unpaid}}


def write_workbook(model: dict, path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    summary = wb.active
    summary.title = "Investor summary"
    summary.append(["Five tier hybrid waterfall", "Value"])
    for key, value in model["summary"].items():
        summary.append([key.replace("_", " ").title(), value if value is not None else "n.a."])
    detail = wb.create_sheet("Period waterfall")
    headers = list(model["periods"][0])
    detail.append([k.replace("_", " ").title() for k in headers])
    for r in model["periods"]:
        detail.append([date.fromisoformat(r[k]) if k == "date" else r[k] for k in headers])
    detail.freeze_panes = "B2"
    for cell in detail["A"][1:]:
        cell.number_format = "mm/dd/yyyy"
    for row in detail.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = '#,##0.00;(#,##0.00);–'
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions if sheet == detail else None
        sheet.row_dimensions[1].height = 26
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="12283F")
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        for col in sheet.columns:
            letter = get_column_letter(col[0].column)
            sheet.column_dimensions[letter].width = 21 if letter != "A" else 29
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    data = json.loads(args.input.read_text())
    model = calculate(data["periods"], Terms(**data.get("terms", {})))
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "waterfall_results.json").write_text(json.dumps(model, indent=2) + "\n")
    write_workbook(model, args.output_dir / "waterfall_results.xlsx")
    print(json.dumps(model["summary"], indent=2))


if __name__ == "__main__":
    main()
